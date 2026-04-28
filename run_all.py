import os

from pandas import ExcelWriter
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import time
import platform
import code_ai.shared_variables as sv
from main import main
from config.load import load_config, load_model_config
import pandas as pd
import gc

def save_metrics_to_excel(
        average_dataset_error, average_model_error,
        dataset_metrics_all, model_metrics_all, hyperparameters,
        model_type, excel_filename="metrics_results.xlsx",
        param_summary=None
):

    # Define formatting styles
    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True)
    center_alignment = Alignment(horizontal="center")
    gray_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    thin_border = Border(left=Side(style='thin'),
                         right=Side(style='thin'),
                         top=Side(style='thin'),
                         bottom=Side(style='thin'))

    # Check if the file exists; if not, create it with initial sheets
    if not os.path.exists(excel_filename):
        with ExcelWriter(excel_filename, engine='openpyxl') as writer:
            pd.DataFrame().to_excel(writer, sheet_name='Dataset Error', index=False)
            pd.DataFrame().to_excel(writer, sheet_name='Model Error', index=False)
            pd.DataFrame().to_excel(writer, sheet_name='Hyperparameters', index=False)

    # Open the file in append mode
    with ExcelWriter(excel_filename, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
        workbook = writer.book
        sheets = ['Dataset Error', 'Model Error']
        avg = [average_dataset_error, average_model_error]
        metrics_all = [dataset_metrics_all, model_metrics_all]
        for n in range(len(sheets)):
        # Access the sheets
            if sheets[n] not in workbook.sheetnames:
                data_sheet = workbook.create_sheet(sheets[n])
            else:
                data_sheet = workbook[sheets[n]]

            # Determine starting column for new model in each sheet
            data_start_col = data_sheet.max_column + 2 if data_sheet.max_column > 1 else 1

            # Write the model name as a header in the first row of the new column
            data_sheet.cell(row=1, column=data_start_col, value=model_type).font = title_font
            data_sheet.cell(row=1, column=data_start_col).alignment = center_alignment

            # Write Average Dataset Error Metrics under the model name
            metrics_headers_dataset = avg[n].columns
            for idx, header in enumerate(metrics_headers_dataset):
                cell = data_sheet.cell(row=2, column=data_start_col + idx, value=header)
                cell.font = header_font
                cell.fill = gray_fill
                cell.border = thin_border

            # Write Average Dataset Error Metrics Data
            for i, row in avg[n].iterrows():
                for j, value in enumerate(metrics_headers_dataset):
                    data_sheet.cell(row=i + 3, column=data_start_col + j, value=row[value]).border = thin_border

            # Write Per-Step Dataset Error Metrics
            per_step_dataset_start_row = avg[n].shape[0] + 5
            data_sheet.cell(row=per_step_dataset_start_row, column=data_start_col, value=" ").font = title_font
            data_sheet.cell(row=per_step_dataset_start_row, column=data_start_col).alignment = center_alignment

            # Define headers for per-step dataset metrics
            per_step_headers_dataset = metrics_all[n].columns
            for idx, header in enumerate(per_step_headers_dataset):
                cell = data_sheet.cell(row=per_step_dataset_start_row + 1, column=data_start_col + idx, value=header)
                cell.font = header_font
                cell.fill = gray_fill
                cell.border = thin_border

            # Write per-step dataset metrics
            current_row = per_step_dataset_start_row + 2
            for _, row in metrics_all[n].iterrows():
                for j, value in enumerate(per_step_headers_dataset):
                    data_sheet.cell(row=current_row, column=data_start_col + j, value=row[value]).border = thin_border
                current_row += 1

            # Adjust column widths for better readability
            def adjust_column_width(sheet):
                for column in sheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if cell.value:
                                max_length = max(max_length, len(str(cell.value)))
                        except:
                            pass
                    adjusted_width = (max_length + 2)
                    sheet.column_dimensions[column_letter].width = adjusted_width

            adjust_column_width(data_sheet)
        if hyperparameters:
            # Write hyperparameters to Hyperparameters sheet
            if 'Hyperparameters' not in workbook.sheetnames:
                hyperparams_sheet = workbook.create_sheet('Hyperparameters')
            else:
                hyperparams_sheet = workbook['Hyperparameters']

            # Determine starting column
            hyper_start_col = hyperparams_sheet.max_column + 2 if hyperparams_sheet.max_column > 1 else 1

            # Model name as header
            hyperparams_sheet.cell(row=1, column=hyper_start_col, value=model_type).font = title_font
            hyperparams_sheet.cell(row=1, column=hyper_start_col).alignment = center_alignment

            row = 2
            # Best Hyperparameters Header
            hyperparams_sheet.cell(row=row, column=hyper_start_col, value="Best Hyperparameters").font = header_font
            row += 1
            hyperparams_sheet.cell(row=row, column=hyper_start_col, value="Parameter").font = header_font
            hyperparams_sheet.cell(row=row, column=hyper_start_col + 1, value="Value").font = header_font
            row += 1

            # Write Best Hyperparameters
            if hyperparameters:
                for param_name, param_value in hyperparameters.items():
                    hyperparams_sheet.cell(row=row, column=hyper_start_col, value=param_name)
                    hyperparams_sheet.cell(row=row, column=hyper_start_col + 1, value=str(param_value))
                    row += 1

            # Add a space
            row += 1

            # Summary Header
            hyperparams_sheet.cell(row=row, column=hyper_start_col, value="Parameter Summary").font = header_font
            row += 1
            hyperparams_sheet.cell(row=row, column=hyper_start_col, value="Parameter").font = header_font
            hyperparams_sheet.cell(row=row, column=hyper_start_col + 1, value="Min Tried").font = header_font
            hyperparams_sheet.cell(row=row, column=hyper_start_col + 2, value="Max Tried").font = header_font
            row += 1

            if param_summary:
                for p_name, p_range in param_summary.items():
                    hyperparams_sheet.cell(row=row, column=hyper_start_col, value=p_name)
                    hyperparams_sheet.cell(row=row, column=hyper_start_col + 1, value=str(p_range['min']))
                    hyperparams_sheet.cell(row=row, column=hyper_start_col + 2, value=str(p_range['max']))
                    row += 1

            # Adjust column width
            adjust_column_width(hyperparams_sheet)

    print(f"Metrics for model {model_type} appended to {excel_filename}")

def clear_screen():
    """
    Clears the terminal screen based on the underlying operating system.
    """
    current_os = platform.system()
    if current_os == "Windows":
        os.system('cls')
    else:
        os.system('clear')

def main_1():
    config = load_config()
    MODEL_CONFIG = load_model_config()
    TIMESTEPS = MODEL_CONFIG['prediction_config']['timesteps']
    STEPS_AHEAD = MODEL_CONFIG['prediction_config']['steps_ahead']

    # Set a single Excel filename for all model results
    # dir = f"metrics_results_ore_lags_{TIMESTEPS}_ahead_{STEPS_AHEAD}_10_min_all"#_dl_corr_ml_all"
    dir = f"metrics_results_ore_lags_{TIMESTEPS}_ahead_{STEPS_AHEAD}_10_min_dl_gate_ml_all"
    sv.OUTPUT_DIR = f"tests/"+dir
    os.makedirs(sv.OUTPUT_DIR, exist_ok=True)
    excel_filename = (sv.OUTPUT_DIR+"/"+dir+".xlsx")
    for model_type in config['ml_model'] + config['deep_model']:
    # for model_type in config['deep_model']:
    # for model_type in config['ml_model']:

    # for model_type in config['ml_model']:
        # Record the start time
        start_time = time.time()
        (average_dataset_error, train_error_all, dataset_metrics_all,
         validation_error_all, best_params, param_summary) = main(model=model_type,
                                                                   feature_filtering_method="Gate", #"ccf" "correlation" "mi" gate
                                                                   size_data = 'All'#'24ME',# 'All'
         )
        # Record the end time
        end_time = time.time()

        elapsed_time = end_time - start_time # Calculate the elapsed time in seconds
        hours, rem = divmod(elapsed_time, 3600)
        minutes, seconds = divmod(rem, 60)
        if best_params:
            param_summary['run_time'] = {'min': str("{:0>2}:{:0>2}:{:05.2f}".format(int(hours), int(minutes), seconds)),
                'max': "", 'type': 'info'}
        else:
            best_params = {'run_time': str("{:0>2}:{:0>2}:{:05.2f}".format(int(hours), int(minutes), seconds))}

        save_metrics_to_excel(
            average_dataset_error=average_dataset_error,
            average_model_error=train_error_all.reset_index(drop=True),
            dataset_metrics_all=dataset_metrics_all,
            model_metrics_all=validation_error_all.reset_index(drop=True),
            model_type=model_type,
            hyperparameters=best_params,
            excel_filename=excel_filename,
            param_summary=param_summary
        )
        clear_screen()
        # Force garbage collection
        gc.collect()

if __name__ == "__main__":
    main_1()